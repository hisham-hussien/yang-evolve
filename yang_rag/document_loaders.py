"""
Document loaders for various file formats
"""
import os
import xml.etree.ElementTree as ET
from typing import Dict, List
from PyPDF2 import PdfReader
from docx import Document
from bs4 import BeautifulSoup
from rich.console import Console

console = Console()


class DocumentLoader:
    """Base class for document loaders"""
    
    def __init__(self):
        pass
    
    def load(self, filepath: str) -> Dict[str, any]:
        """Load a document and return its content and metadata"""
        raise NotImplementedError


class XMLLoader(DocumentLoader):
    """Loader for XML files"""
    
    def load(self, filepath: str) -> Dict[str, any]:
        try:
            tree = ET.parse(filepath)
            root = tree.getroot()
            
            # Extract text content from XML
            text_content = self._extract_text_from_element(root)
            
            return {
                'content': text_content,
                'metadata': {
                    'source': filepath,
                    'file_type': 'xml',
                    'root_tag': root.tag
                }
            }
        except Exception as e:
            console.print(f"[red]Error loading XML file {filepath}: {e}[/red]")
            return None
    
    def _extract_text_from_element(self, element, level=0) -> str:
        """Recursively extract text from XML elements"""
        text_parts = []
        
        # Add element tag as header
        if level < 3:  # Only add headers for top 3 levels
            text_parts.append(f"\n{'#' * (level + 1)} {element.tag}\n")
        
        # Add element text
        if element.text and element.text.strip():
            text_parts.append(element.text.strip())
        
        # Add attributes
        if element.attrib:
            attrs = ", ".join([f"{k}={v}" for k, v in element.attrib.items()])
            text_parts.append(f"[Attributes: {attrs}]")
        
        # Process child elements
        for child in element:
            text_parts.append(self._extract_text_from_element(child, level + 1))
            if child.tail and child.tail.strip():
                text_parts.append(child.tail.strip())
        
        return "\n".join(text_parts)


class PDFLoader(DocumentLoader):
    """Loader for PDF files"""
    
    def load(self, filepath: str) -> Dict[str, any]:
        try:
            reader = PdfReader(filepath)
            text_content = []
            
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_content.append(f"\n--- Page {page_num + 1} ---\n{text}")
            
            return {
                'content': "\n".join(text_content),
                'metadata': {
                    'source': filepath,
                    'file_type': 'pdf',
                    'num_pages': len(reader.pages)
                }
            }
        except Exception as e:
            console.print(f"[red]Error loading PDF file {filepath}: {e}[/red]")
            return None


class DOCXLoader(DocumentLoader):
    """Loader for DOCX files"""
    
    def load(self, filepath: str) -> Dict[str, any]:
        try:
            doc = Document(filepath)
            text_content = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)
            
            return {
                'content': "\n".join(text_content),
                'metadata': {
                    'source': filepath,
                    'file_type': 'docx',
                    'num_paragraphs': len(doc.paragraphs)
                }
            }
        except Exception as e:
            console.print(f"[red]Error loading DOCX file {filepath}: {e}[/red]")
            return None


class TextLoader(DocumentLoader):
    """Loader for text and markdown files"""
    
    def load(self, filepath: str) -> Dict[str, any]:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                content = f.read()
            
            return {
                'content': content,
                'metadata': {
                    'source': filepath,
                    'file_type': 'text'
                }
            }
        except Exception as e:
            console.print(f"[red]Error loading text file {filepath}: {e}[/red]")
            return None


class CodeLoader(DocumentLoader):
    """Loader for code files"""
    
    def load(self, filepath: str) -> Dict[str, any]:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Get file extension to determine language
            _, ext = os.path.splitext(filepath)
            
            # Add context about the code file
            formatted_content = f"""# Code File: {os.path.basename(filepath)}
Language: {ext[1:] if ext else 'unknown'}

```{ext[1:] if ext else ''}
{content}
```
"""
            
            return {
                'content': formatted_content,
                'metadata': {
                    'source': filepath,
                    'file_type': 'code',
                    'language': ext[1:] if ext else 'unknown'
                }
            }
        except Exception as e:
            console.print(f"[red]Error loading code file {filepath}: {e}[/red]")
            return None


def get_loader_for_file(filepath: str) -> DocumentLoader:
    """Get the appropriate loader based on file extension"""
    _, ext = os.path.splitext(filepath)
    ext = ext.lower()
    
    if ext == '.xml':
        return XMLLoader()
    elif ext == '.pdf':
        return PDFLoader()
    elif ext in ['.docx', '.doc']:
        return DOCXLoader()
    elif ext in ['.txt', '.md', '.markdown']:
        return TextLoader()
    elif ext in ['.py', '.js', '.ts', '.java', '.cpp', '.c', '.cs', '.go', '.rs', '.rb', '.php', '.swift', '.kt', '.scala']:
        return CodeLoader()
    else:
        # Default to text loader
        return TextLoader()


def load_document(filepath: str) -> Dict[str, any]:
    """Load a document using the appropriate loader"""
    loader = get_loader_for_file(filepath)
    return loader.load(filepath)
